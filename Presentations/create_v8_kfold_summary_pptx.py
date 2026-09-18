from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
RUN_DIR = PROJECT_ROOT / "analysis_data_drift" / "kfold_runs" / "2026-09-17_161605_V8_kfold_overnight"
RESULTS_CSV = RUN_DIR / "kfold_results.csv"
OUTPUT_PPTX = Path(__file__).resolve().parent / "V8_KFold_Overnight_Summary_2026-09-18.pptx"


DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
MED_BLUE = RGBColor(0x2E, 0x5E, 0x8E)
LIGHT_BLUE = RGBColor(0xD6, 0xE8, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x58, 0x58, 0x58)
LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)
GREEN = RGBColor(0x2D, 0x8B, 0x4E)
RED = RGBColor(0xC0, 0x39, 0x2B)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)


def add_title_bar(slide, title: str) -> None:
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.85))
    rect.fill.solid()
    rect.fill.fore_color.rgb = DARK_BLUE
    rect.line.fill.background()
    tf = rect.text_frame
    tf.margin_left = Inches(0.4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = WHITE


def add_textbox(slide, left, top, width, height, text, font_size=16, bold=False,
                color=BLACK, align=PP_ALIGN.LEFT) -> None:
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color


def add_metric_box(slide, left, top, width, height, title, value, subtitle, fill_color) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = title
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = value
    r2.font.size = Pt(24)
    r2.font.bold = True
    r2.font.color.rgb = WHITE

    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run()
    r3.text = subtitle
    r3.font.size = Pt(11)
    r3.font.color.rgb = WHITE


def set_cell(cell, text, font_size=11, bold=False, font_color=BLACK, fill_color=None, align=PP_ALIGN.CENTER):
    cell.text = str(text)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    run = p.runs[0] if p.runs else p.add_run()
    run.text = str(text)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = font_color
    if fill_color is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill_color


def add_table(slide, df: pd.DataFrame, left, top, width, height) -> None:
    rows = len(df) + 1
    cols = len(df.columns)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table

    for c, header in enumerate(df.columns):
        set_cell(table.cell(0, c), header, font_size=10, bold=True, font_color=WHITE, fill_color=DARK_BLUE)

    for r, row in enumerate(df.itertuples(index=False), start=1):
        bg = LIGHT_GRAY if r % 2 == 1 else WHITE
        for c, value in enumerate(row):
            set_cell(table.cell(r, c), value, font_size=10, fill_color=bg)


def main() -> None:
    results = pd.read_csv(RESULTS_CSV)

    da_acc = results["da_accuracy"].mean()
    ml_acc = results["ml_accuracy"].mean()
    da_f1 = results["da_f1_macro"].mean()
    ml_f1 = results["ml_f1_macro"].mean()
    override_prec = results["override_precision"].mean()
    net_benefit = results["net_benefit"].mean()
    negative_folds = int((results["net_benefit"] < 0).sum())
    positive_folds = int((results["net_benefit"] > 0).sum())
    results["f1_delta"] = results["ml_f1_macro"] - results["da_f1_macro"]
    results["acc_delta"] = results["ml_accuracy"] - results["da_accuracy"]

    fold5 = results.loc[results["fold"] == 5].iloc[0]

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "V8 5-Fold Validation: Current ML Is Not Ready for Reliable Deployment")

    add_textbox(
        slide,
        Inches(0.6), Inches(1.0), Inches(12.1), Inches(0.8),
        "Study-level 5-fold validation across all 183 available files shows that the current V8 impedance-only model does not generalize reliably enough to replace or override DA.",
        font_size=18, bold=False, color=GRAY,
    )

    add_metric_box(slide, Inches(0.7), Inches(1.95), Inches(2.3), Inches(1.35), "DA Macro F1", f"{da_f1:.4f}", "mean across 5 outer folds", DARK_BLUE)
    add_metric_box(slide, Inches(3.15), Inches(1.95), Inches(2.3), Inches(1.35), "ML Macro F1", f"{ml_f1:.4f}", "mean across 5 outer folds", MED_BLUE)
    add_metric_box(slide, Inches(5.60), Inches(1.95), Inches(2.3), Inches(1.35), "Override Precision", f"{override_prec:.4f}", "target was much higher", ORANGE)
    add_metric_box(slide, Inches(8.05), Inches(1.95), Inches(2.3), Inches(1.35), "Mean Net Benefit", f"{net_benefit:+.0f}", "samples per fold", RED if net_benefit < 0 else GREEN)
    add_metric_box(slide, Inches(10.50), Inches(1.95), Inches(2.1), Inches(1.35), "Negative Folds", f"{negative_folds}/5", "only one fold was positive", RED)

    summary_text = (
        f"DA accuracy {da_acc:.4f} vs ML accuracy {ml_acc:.4f}\n"
        f"DA macro F1 {da_f1:.4f} vs ML macro F1 {ml_f1:.4f}\n"
        f"Correct overrides 285,331 vs harmful overrides 420,429 across the full benchmark\n"
        "Conclusion: this V8 configuration is not robust enough to lock as a production override model."
    )
    add_textbox(slide, Inches(0.8), Inches(3.55), Inches(11.8), Inches(1.35), summary_text, font_size=17)

    actions = (
        "Recommended next steps\n"
        "1. Do not search for a lucky split; this result already answers the generalization question.\n"
        "2. Reframe ML as assistive only, or add orthogonal signals such as pressure/contact force.\n"
        "3. Use this as the management conclusion: ML capability exists in some cohorts, but current V8 is not deployment-ready."
    )
    add_textbox(slide, Inches(0.8), Inches(5.05), Inches(11.8), Inches(1.65), actions, font_size=16)

    footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.95), Inches(12.0), Inches(0.3))
    footer.text_frame.text = "Source: analysis_data_drift/kfold_runs/2026-09-17_161605_V8_kfold_overnight"
    footer.text_frame.paragraphs[0].font.size = Pt(10)
    footer.text_frame.paragraphs[0].font.color.rgb = GRAY

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "Fold 5 Was Slightly Positive, But It Does Not Change the Overall Conclusion")

    table_df = results[["fold", "n_train_files", "n_test_files", "da_f1_macro", "ml_f1_macro", "override_precision", "net_benefit"]].copy()
    table_df["da_f1_macro"] = table_df["da_f1_macro"].map(lambda x: f"{x:.4f}")
    table_df["ml_f1_macro"] = table_df["ml_f1_macro"].map(lambda x: f"{x:.4f}")
    table_df["override_precision"] = table_df["override_precision"].map(lambda x: f"{x:.4f}")
    table_df["net_benefit"] = table_df["net_benefit"].map(lambda x: f"{x:+.0f}")
    table_df.columns = ["Fold", "Train", "Test", "DA F1", "ML F1", "Override Prec", "Net Benefit"]
    add_table(slide, table_df, Inches(0.55), Inches(1.15), Inches(6.6), Inches(4.55))

    fold5_text = (
        f"Fold 5 details\n"
        f"Net benefit: {int(fold5['net_benefit']):+d}\n"
        f"Override precision: {fold5['override_precision']:.4f}\n"
        f"DA F1: {fold5['da_f1_macro']:.4f}\n"
        f"ML F1: {fold5['ml_f1_macro']:.4f}"
    )
    add_textbox(slide, Inches(7.45), Inches(1.2), Inches(2.0), Inches(1.6), fold5_text, font_size=15, bold=False)

    interpret = (
        "Interpretation\n"
        "Fold 5 was only mildly positive. DA still had slightly better macro F1, but the override decisions did slightly more good than harm in that cohort.\n"
        "This is evidence that ML can sometimes help, not evidence that the current model is reliable.\n"
        "Four other outer folds were negative, including one strongly negative fold at -88,330."
    )
    add_textbox(slide, Inches(7.45), Inches(3.0), Inches(5.2), Inches(1.95), interpret, font_size=15)

    ten_fold = (
        "Would 10-fold help?\n"
        "Probably not enough to change the conclusion. On 183 studies, 10-fold would use only 18 to 19 test studies per fold instead of 36 to 37.\n"
        "That makes each fold noisier and nearly doubles runtime, while the current 5-fold result is already consistently unfavorable overall.\n"
        "10-fold could be used later for a finer estimate, but 5-fold is the better practical choice here."
    )
    add_textbox(slide, Inches(7.45), Inches(5.15), Inches(5.2), Inches(1.75), ten_fold, font_size=14)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, "Technical View: The Model Is Close on Classification but Fails on Safe Overrides")

    tech_intro = (
        "Across all 5 outer folds, ML classification stayed close to DA, but it was almost always slightly worse. "
        "The larger failure mode was override safety: too many harmful overrides relative to correct ones."
    )
    add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12.0), Inches(0.75), tech_intro, font_size=17, color=GRAY)

    add_metric_box(slide, Inches(0.7), Inches(1.9), Inches(2.2), Inches(1.25), "Positive Folds", f"{positive_folds}", "out of 5 outer folds", GREEN if positive_folds > 0 else RED)
    add_metric_box(slide, Inches(3.05), Inches(1.9), Inches(2.4), Inches(1.25), "Mean F1 Delta", f"{results['f1_delta'].mean():+.4f}", "ML minus DA", RED if results['f1_delta'].mean() < 0 else GREEN)
    add_metric_box(slide, Inches(5.6), Inches(1.9), Inches(2.4), Inches(1.25), "Mean Acc Delta", f"{results['acc_delta'].mean():+.4f}", "ML minus DA", RED if results['acc_delta'].mean() < 0 else GREEN)
    add_metric_box(slide, Inches(8.15), Inches(1.9), Inches(2.15), Inches(1.25), "Correct Overrides", "285,331", "summed across folds", MED_BLUE)
    add_metric_box(slide, Inches(10.45), Inches(1.9), Inches(2.15), Inches(1.25), "Harmful Overrides", "420,429", "summed across folds", RED)

    tech_table = results[["fold", "n_test_files", "acc_delta", "f1_delta", "override_precision", "net_benefit"]].copy()
    tech_table["acc_delta"] = tech_table["acc_delta"].map(lambda x: f"{x:+.4f}")
    tech_table["f1_delta"] = tech_table["f1_delta"].map(lambda x: f"{x:+.4f}")
    tech_table["override_precision"] = tech_table["override_precision"].map(lambda x: f"{x:.4f}")
    tech_table["net_benefit"] = tech_table["net_benefit"].map(lambda x: f"{x:+.0f}")
    tech_table.columns = ["Fold", "Test N", "Acc Delta", "F1 Delta", "Override Prec", "Net Benefit"]
    add_table(slide, tech_table, Inches(0.7), Inches(3.5), Inches(6.3), Inches(3.0))

    technical_points = (
        "Technical interpretation\n"
        "1. ML minus DA F1 was negative in every outer fold, ranging from -0.0043 to -0.0410.\n"
        "2. Accuracy was also worse in 4 of 5 folds, with only a trivial +0.0005 gain in fold 5.\n"
        "3. Override precision varied from 0.2531 to 0.4423, never reaching a clinically safe regime.\n"
        "4. The model is learning something real, but the current impedance-only representation is not stable enough across cohorts to support production override behavior."
    )
    add_textbox(slide, Inches(7.35), Inches(3.45), Inches(5.35), Inches(2.7), technical_points, font_size=14)

    closing = (
        "Most defensible technical conclusion\n"
        "The present V8 model is not a deployment candidate. Its best near-term value would be as a research benchmark, "
        "an uncertainty/assistive signal, or a baseline to compare against future pressure-augmented models."
    )
    add_textbox(slide, Inches(7.35), Inches(6.1), Inches(5.35), Inches(0.9), closing, font_size=13)

    prs.save(OUTPUT_PPTX)
    print(f"Saved presentation -> {OUTPUT_PPTX}")


if __name__ == "__main__":
    main()