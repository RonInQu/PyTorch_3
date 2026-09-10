"""
Generate Word documents from Invention Disclosure markdown files.

Equations marked with $$...$$ are converted to native Word OMML equations
using the pipeline:
    LaTeX  →  MathML  (via latex2mathml)
              MathML  →  OMML   (via Microsoft's MML2OMML.XSL + lxml XSLT)
                         OMML  injected as Word-native math objects

Inline $...$ expressions are also rendered to OMML and inserted inline.
No PNG images are used.
"""

import re
import io
import sys
from pathlib import Path

import lxml.etree as etree
from latex2mathml import converter as l2m
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

# ── XSLT stylesheet (Microsoft MML2OMML.XSL ships with Office) ───────────────
MML2OMML_XSL = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS    = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

_xslt_transform = None


def _get_transform():
    global _xslt_transform
    if _xslt_transform is None:
        xslt_tree = etree.parse(str(MML2OMML_XSL))
        _xslt_transform = etree.XSLT(xslt_tree)
    return _xslt_transform


def latex_to_omml_str(latex: str) -> str:
    """Convert LaTeX (no delimiters) to an OMML XML string (<m:oMath ...>)."""
    latex = latex.strip()
    mathml_str = l2m.convert(latex, xmlns="http://www.w3.org/1998/Math/MathML")
    mml_tree = etree.fromstring(mathml_str.encode())
    omml_tree = _get_transform()(mml_tree)
    root = omml_tree.getroot()
    return etree.tostring(root, encoding="unicode")


# ── Helpers to insert equations into the Word document ───────────────────────

def add_block_equation(doc: Document, latex: str):
    """Insert a centred display equation as a native Word OMML paragraph."""
    try:
        omml_str = latex_to_omml_str(latex)
        # Build a full <w:p> containing <m:oMathPara><m:oMath>…</m:oMath></m:oMathPara>
        # with explicit namespace declarations so parse_xml knows every prefix.
        p_xml = (
            f'<w:p xmlns:w="{W_NS}" xmlns:m="{MATH_NS}">'
            f'<w:pPr><w:jc w:val="center"/></w:pPr>'
            f'<m:oMathPara xmlns:m="{MATH_NS}">'
            f'{omml_str}'
            f'</m:oMathPara>'
            f'</w:p>'
        )
        p_el = parse_xml(p_xml)
        # Insert before the section-properties element (end of body)
        body = doc.element.body
        sect_pr = body.find(qn('w:sectPr'))
        if sect_pr is not None:
            body.insert(list(body).index(sect_pr), p_el)
        else:
            body.append(p_el)
    except Exception as exc:
        print(f"  [OMML fallback] {latex[:60]!r}: {exc}")
        fb = doc.add_paragraph(f"  {latex}")
        fb.paragraph_format.left_indent = Inches(0.5)


def _build_inline_omml_el(latex: str):
    """Return a parsed <m:oMath> lxml element for inline insertion."""
    omml_str = latex_to_omml_str(latex)
    return parse_xml(f'<m:oMath xmlns:m="{MATH_NS}">{omml_str[omml_str.index(">")+1:omml_str.rfind("<")]}</m:oMath>')


def add_paragraph_with_inline_math(doc: Document, text: str, style: str = "Normal"):
    """
    Add a paragraph containing mixed plain text and inline $...$ equations.
    Equations become native <m:oMath> objects; surrounding text are <w:r> runs.
    """
    # Build the full paragraph XML manually so all pieces share one namespace context.
    runs_xml = ""
    parts = re.split(r'(\$[^$]+\$)', text)
    for part in parts:
        m = re.fullmatch(r'\$([^$]+)\$', part)
        if m:
            try:
                omml_str = latex_to_omml_str(m.group(1))
                runs_xml += omml_str
            except Exception as exc:
                print(f"  [inline fallback] {m.group(1)!r}: {exc}")
                clean = re.sub(r'\*\*(.+?)\*\*', r'\1', m.group(1))
                runs_xml += f'<w:r xmlns:w="{W_NS}"><w:t>{clean}</w:t></w:r>'
        else:
            if part:
                clean = re.sub(r'\*\*(.+?)\*\*', r'\1', part)
                clean = re.sub(r'\*(.+?)\*', r'\1', clean)
                # Escape XML special chars
                clean = clean.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                runs_xml += (
                    f'<w:r xmlns:w="{W_NS}">'
                    f'<w:t xml:space="preserve">{clean}</w:t>'
                    f'</w:r>'
                )
    p_xml = (
        f'<w:p xmlns:w="{W_NS}" xmlns:m="{MATH_NS}">'
        f'{runs_xml}'
        f'</w:p>'
    )
    p_el = parse_xml(p_xml)
    body = doc.element.body
    sect_pr = body.find(qn('w:sectPr'))
    if sect_pr is not None:
        body.insert(list(body).index(sect_pr), p_el)
    else:
        body.append(p_el)
    # Apply style via a thin wrapper paragraph lookup
    if style != "Normal":
        try:
            style_id = doc.styles[style].style_id
            pPr = p_el.find(qn('w:pPr'))
            if pPr is None:
                pPr = etree.SubElement(p_el, qn('w:pPr'))
                p_el.insert(0, pPr)
            pStyle = etree.SubElement(pPr, qn('w:pStyle'))
            pStyle.set(qn('w:val'), style_id)
        except Exception:
            pass
    return p_el


# ── Table builder ─────────────────────────────────────────────────────────────

def add_markdown_table(doc: Document, lines_block: list):
    rows_data = []
    for line in lines_block:
        line = line.strip()
        if not line or re.match(r'^\|[-| :]+\|$', line):
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        rows_data.append(cells)
    if not rows_data:
        return
    ncols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=ncols)
    table.style = 'Medium Shading 1 Accent 1'
    for i, row_data in enumerate(rows_data):
        for j, cell_text in enumerate(row_data):
            if j < ncols:
                table.rows[i].cells[j].text = cell_text
    doc.add_paragraph('')


# ── Main document builder ─────────────────────────────────────────────────────

def build_docx(md_path: Path, out_path: Path):
    lines = md_path.read_text(encoding='utf-8').splitlines()
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        s = lines[i].rstrip()

        # ── Block equation  $$...$$  (possibly multi-line) ───────────────────
        if s.strip().startswith('$$'):
            eq_lines = [s.strip().lstrip('$')]
            if s.strip() == '$$' or not s.strip().rstrip('$') or s.strip() == s.strip().lstrip('$') + '$$':
                # opening $$ on its own line — collect until closing $$
                eq_lines = []
                i += 1
                while i < len(lines):
                    row = lines[i].rstrip()
                    if row.strip() == '$$':
                        break
                    eq_lines.append(row)
                    if row.strip().endswith('$$') and row.strip() != '$$':
                        eq_lines[-1] = row.strip().rstrip('$').rstrip()
                        break
                    i += 1
            else:
                # inline: $$expr$$
                eq_lines = [s.strip().strip('$')]

            latex = ' '.join(eq_lines).strip().strip('$').strip()
            add_block_equation(doc, latex)
            i += 1
            continue

        # ── Markdown table ───────────────────────────────────────────────────
        if s.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].startswith('|'):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        # ── Blank line ────────────────────────────────────────────────────────
        if not s:
            doc.add_paragraph('')
            i += 1
            continue

        # ── Headings ──────────────────────────────────────────────────────────
        if s.startswith('#### '):
            doc.add_heading(s[5:].strip(), level=4)
        elif s.startswith('### '):
            doc.add_heading(s[4:].strip(), level=3)
        elif s.startswith('## '):
            doc.add_heading(s[3:].strip(), level=2)
        elif s.startswith('# '):
            doc.add_heading(s[2:].strip(), level=1)

        # ── Numbered list ─────────────────────────────────────────────────────
        elif re.match(r'^\d+\.\s+', s):
            body = re.sub(r'^\d+\.\s+', '', s)
            if '$' in body:
                add_paragraph_with_inline_math(doc, body, style='List Number')
            else:
                doc.add_paragraph(body, style='List Number')

        # ── Bullet list ───────────────────────────────────────────────────────
        elif s.startswith('- '):
            body = s[2:].strip()
            if '$' in body:
                add_paragraph_with_inline_math(doc, body, style='List Bullet')
            else:
                doc.add_paragraph(body, style='List Bullet')

        # ── Body paragraph (may contain inline math) ──────────────────────────
        else:
            if '$' in s:
                add_paragraph_with_inline_math(doc, s)
            else:
                clean = re.sub(r'\*\*(.+?)\*\*', r'\1', s)
                clean = re.sub(r'\*(.+?)\*', r'\1', clean)
                doc.add_paragraph(clean)

        i += 1

    doc.save(out_path)
    print(f'Saved: {out_path}')


# ── Entry point ───────────────────────────────────────────────────────────────

if __name__ == '__main__':
    base = Path(__file__).parent

    files = [
        ('Tetrapolar_Electrode_Impedance_Invention_Disclosure.md',
         'Tetrapolar_Electrode_Impedance_Invention_Disclosure.docx'),
        ('Tetrapolar_Electrode_Impedance_Invention_Disclosure_PatentCounsel.md',
         'Tetrapolar_Electrode_Impedance_Invention_Disclosure_PatentCounsel.docx'),
    ]

    for md_name, docx_name in files:
        md_path = base / md_name
        out_path = base / docx_name
        if not md_path.exists():
            print(f'Skipping (not found): {md_path}')
            continue
        print(f'Building {docx_name} ...')
        build_docx(md_path, out_path)

    print('Done.')
